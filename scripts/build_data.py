#!/usr/bin/env python3
"""Gera public/bpmn/*.bpmn e public/data/*.json a partir dos arquivos originais em /source.
Nada é digitado à mão: todos os números vêm da planilha e todos os fluxos dos .bpm (Bizagi)."""
import zipfile, io, json, re, statistics, datetime, os, sys
import xml.etree.ElementTree as ET
from xml.sax.saxutils import escape, quoteattr
import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(ROOT, 'source'); OUT_B = os.path.join(ROOT, 'public/bpmn'); OUT_D = os.path.join(ROOT, 'public/data')
os.makedirs(OUT_B, exist_ok=True); os.makedirs(OUT_D, exist_ok=True)
NS = '{http://www.wfmc.org/2009/XPDL2.2}'
Q = lambda t: NS + t
clean = lambda s: re.sub(r'\s+', ' ', (s or '').replace('\xa0', ' ')).strip()

# ---------------------------------------------------------------- BPMN
PROCS = [
  dict(id='glaucoma', file='Fluxograma_Glaucoma.bpm', title='Glaucoma', short='Jornada do paciente com glaucoma'),
  dict(id='exame-de-cornea', file='Diagrama_Cornea.bpmn', title='Exame de córnea', short='Exame de córnea (paciente nova)'),
  dict(id='teste-de-lente', file='Diagrama_Teste_de_lente.bpmn', title='Teste de lente', short='Consulta de teste de lente'),
]

def read_diagram(path):
    """devolve o Diagram.xml (XPDL) que contém um processo com atividades"""
    outer = zipfile.ZipFile(path)
    best = None
    for n in outer.namelist():
        if not n.endswith('.diag'): continue
        inner = zipfile.ZipFile(io.BytesIO(outer.read(n)))
        root = ET.fromstring(inner.read('Diagram.xml'))
        n_act = len(list(root.iter(Q('Activity'))))
        if best is None or n_act > best[0]: best = (n_act, root)
    return best[1]

def argb(v):
    v = int(v) & 0xFFFFFFFF
    return '#%02x%02x%02x' % ((v >> 16) & 255, (v >> 8) & 255, v & 255)


def via_map(nodes, flows):
    """rótulo (Sim/Não…) do fluxo que sai de uma decisão e chega ao nó"""
    kinds = {n['id']: n['kind'] for n in nodes}; v = {}
    for f in flows:
        if kinds.get(f['src']) == 'exclusiveGateway' and f['name']: v[f['dst']] = f['name']
    return v


def seq_map(nodes, flows):
    """ordem causal do percurso: ordenação topológica pelos fluxos (laços de volta ignorados), desempate pela posição x no desenho"""
    import heapq
    byid = {n['id']: n for n in nodes}; succ = {n['id']: [] for n in nodes}
    for f in flows: succ[f['src']].append(f['dst'])
    back = set(); state = {}
    def dfs(u):
        state[u] = 1
        for v in succ[u]:
            if state.get(v) == 1: back.add((u, v))
            elif v not in state: dfs(v)
        state[u] = 2
    for n in sorted(nodes, key=lambda n: n['x']):
        if n['kind'] == 'startEvent' and n['id'] not in state: dfs(n['id'])
    for n in nodes:
        if n['id'] not in state: dfs(n['id'])
    indeg = {n['id']: 0 for n in nodes}
    for f in flows:
        if (f['src'], f['dst']) not in back: indeg[f['dst']] += 1
    heap = [(byid[i]['x'], byid[i]['y'], i) for i, d in indeg.items() if d == 0]; heapq.heapify(heap); order = {}
    while heap:
        _, _, u = heapq.heappop(heap); order[u] = len(order)
        for v in succ[u]:
            if (u, v) in back: continue
            indeg[v] -= 1
            if indeg[v] == 0: heapq.heappush(heap, (byid[v]['x'], byid[v]['y'], v))
    return order

def node_out(n, via, seq):
    return dict(id=n['id'], name=n['name'], kind=n['kind'], lane=n['lane'], doc=n['doc'], gw=n['gw'], x=round(n['x']), y=round(n['y']), via=via.get(n['id']), seq=seq.get(n['id']))

def convert(proc):
    root = read_diagram(os.path.join(SRC, proc['file']))
    # pool real = a que tem lanes ou atividades
    wp = [w for w in root.iter(Q('WorkflowProcess')) if list(w.iter(Q('Activity')))][0]
    pool = [p for p in root.find(Q('Pools')) if p.get('Process') == wp.get('Id')][0]
    pg = pool.find(Q('NodeGraphicsInfos'))[0]; pc = pg.find(Q('Coordinates'))
    px, py, pw, ph = float(pc.get('XCoordinate')), float(pc.get('YCoordinate')), float(pg.get('Width')), float(pg.get('Height'))
    lanes = []
    for l in pool.find(Q('Lanes')):
        g = l.find('.//' + Q('NodeGraphicsInfo')); c = g.find(Q('Coordinates'))
        lanes.append(dict(id='Lane_' + l.get('Id')[:8], name=clean(l.get('Name')),
                          x=px + float(c.get('XCoordinate')), y=py + float(c.get('YCoordinate')),
                          w=float(g.get('Width')), h=float(g.get('Height'))))
    nodes = []
    for a in wp.iter(Q('Activity')):
        g = a.find('.//' + Q('NodeGraphicsInfo')); c = g.find(Q('Coordinates'))
        x, y, w, h = float(c.get('XCoordinate')), float(c.get('YCoordinate')), float(g.get('Width')), float(g.get('Height'))
        ev, rt = a.find(Q('Event')), a.find(Q('Route'))
        if ev is not None:
            kind = 'startEvent' if ev.find(Q('StartEvent')) is not None else 'endEvent'
        elif rt is not None:
            kind = 'exclusiveGateway'
        else:
            kind = 'task'
        doc = a.find(Q('Documentation'))
        nodes.append(dict(id='Node_' + a.get('Id')[:8], guid=a.get('Id'), name=clean(a.get('Name')), kind=kind,
                          x=x, y=y, w=w, h=h, doc=clean(doc.text) if doc is not None and doc.text else '',
                          gw=(rt.get('GatewayDirection') if rt is not None else None)))
    for n in nodes:
        cy = n['y'] + n['h'] / 2
        n['lane'] = next((l['id'] for l in lanes if l['y'] <= cy <= l['y'] + l['h']), None)
    flows = []
    for t in wp.iter(Q('Transition')):
        cg = t.find('.//' + Q('ConnectorGraphicsInfo'))
        pts = [(float(c.get('XCoordinate')), float(c.get('YCoordinate'))) for c in cg.iter(Q('Coordinates'))]
        flows.append(dict(id='Flow_' + t.get('Id')[:8], src='Node_' + t.get('From')[:8], dst='Node_' + t.get('To')[:8],
                          name=clean(t.get('Name')), pts=pts))
    ids = {n['id'] for n in nodes}
    assert all(f['src'] in ids and f['dst'] in ids for f in flows), 'transição aponta para nó inexistente'
    inc = {n['id']: [] for n in nodes}; out = {n['id']: [] for n in nodes}
    for f in flows: out[f['src']].append(f['id']); inc[f['dst']].append(f['id'])
    pool_name = clean(pool.get('Name'))
    P = []
    P.append('<?xml version="1.0" encoding="UTF-8"?>')
    P.append('<bpmn:definitions xmlns:bpmn="http://www.omg.org/spec/BPMN/20100524/MODEL" xmlns:bpmndi="http://www.omg.org/spec/BPMN/20100524/DI" xmlns:dc="http://www.omg.org/spec/DD/20100524/DC" xmlns:di="http://www.omg.org/spec/DD/20100524/DI" id="Defs_%s" targetNamespace="http://skema.jr/rede-oftalmo">' % proc['id'])
    P.append('<bpmn:collaboration id="Collab_1"><bpmn:participant id="Pool_1" name=%s processRef="Process_1" /></bpmn:collaboration>' % quoteattr(pool_name))
    P.append('<bpmn:process id="Process_1" isExecutable="false">')
    if lanes:
        P.append('<bpmn:laneSet id="LaneSet_1">')
        for l in lanes:
            P.append('<bpmn:lane id="%s" name=%s>' % (l['id'], quoteattr(l['name'])))
            for n in nodes:
                if n['lane'] == l['id']: P.append('<bpmn:flowNodeRef>%s</bpmn:flowNodeRef>' % n['id'])
            P.append('</bpmn:lane>')
        P.append('</bpmn:laneSet>')
    for n in nodes:
        nm = (' name=%s' % quoteattr(n['name'])) if n['name'] else ''
        P.append('<bpmn:%s id="%s"%s>' % (n['kind'], n['id'], nm))
        if n['doc']: P.append('<bpmn:documentation>%s</bpmn:documentation>' % escape(n['doc']))
        for i in inc[n['id']]: P.append('<bpmn:incoming>%s</bpmn:incoming>' % i)
        for o in out[n['id']]: P.append('<bpmn:outgoing>%s</bpmn:outgoing>' % o)
        P.append('</bpmn:%s>' % n['kind'])
    for f in flows:
        nm = (' name=%s' % quoteattr(f['name'])) if f['name'] else ''
        P.append('<bpmn:sequenceFlow id="%s" sourceRef="%s" targetRef="%s"%s />' % (f['id'], f['src'], f['dst'], nm))
    P.append('</bpmn:process>')
    P.append('<bpmndi:BPMNDiagram id="Diagram_1"><bpmndi:BPMNPlane id="Plane_1" bpmnElement="Collab_1">')
    P.append('<bpmndi:BPMNShape id="Pool_1_di" bpmnElement="Pool_1" isHorizontal="true"><dc:Bounds x="%g" y="%g" width="%g" height="%g" /></bpmndi:BPMNShape>' % (px, py, pw, ph))
    for l in lanes:
        P.append('<bpmndi:BPMNShape id="%s_di" bpmnElement="%s" isHorizontal="true"><dc:Bounds x="%g" y="%g" width="%g" height="%g" /></bpmndi:BPMNShape>' % (l['id'], l['id'], l['x'], l['y'], l['w'], l['h']))
    for n in nodes:
        mk = ' isMarkerVisible="true"' if n['kind'] == 'exclusiveGateway' else ''
        P.append('<bpmndi:BPMNShape id="%s_di" bpmnElement="%s"%s><dc:Bounds x="%g" y="%g" width="%g" height="%g" /></bpmndi:BPMNShape>' % (n['id'], n['id'], mk, n['x'], n['y'], n['w'], n['h']))
    for f in flows:
        P.append('<bpmndi:BPMNEdge id="%s_di" bpmnElement="%s">' % (f['id'], f['id']))
        for (x, y) in f['pts']: P.append('<di:waypoint x="%g" y="%g" />' % (x, y))
        P.append('</bpmndi:BPMNEdge>')
    P.append('</bpmndi:BPMNPlane></bpmndi:BPMNDiagram></bpmn:definitions>')
    with open(os.path.join(OUT_B, proc['id'] + '.bpmn'), 'w', encoding='utf-8') as fh: fh.write('\n'.join(P))
    return dict(id=proc['id'], title=proc['title'], short=proc['short'], poolName=pool_name, sourceFile=proc['file'],
                bpmn='bpmn/%s.bpmn' % proc['id'],
                lanes=[dict(id=l['id'], name=l['name']) for l in lanes],
                nodes=[node_out(n, via_map(nodes, flows), seq_map(nodes, flows)) for n in nodes],
                flows=[dict(id=f['id'], src=f['src'], dst=f['dst'], name=f['name']) for f in flows],
                counts=dict(tasks=sum(n['kind'] == 'task' for n in nodes), gateways=sum(n['kind'] == 'exclusiveGateway' for n in nodes),
                            events=sum(n['kind'].endswith('Event') for n in nodes), flows=len(flows), lanes=len(lanes)))


# ---- importador de BPMN 2.0 (export do Bizagi) -> BPMN limpo para o bpmn-js
BM = '{http://www.omg.org/spec/BPMN/20100524/MODEL}'
BD = '{http://www.omg.org/spec/BPMN/20100524/DI}'
DC = '{http://www.omg.org/spec/DD/20100524/DC}'
DI_ = '{http://www.omg.org/spec/DD/20100524/DI}'

def convert_bpmn(proc):
    root = ET.parse(os.path.join(SRC, proc['file'])).getroot()
    coll = root.find(BM + 'collaboration')
    part = next(p for p in coll.findall(BM + 'participant') if p.get('name'))   # a outra participante é vazia
    pr = next(p for p in root.findall(BM + 'process') if p.get('id') == part.get('processRef'))
    pool_id, pool_name = part.get('id'), clean(part.get('name'))
    shapes = {}; edges = {}
    for sh in root.iter(BD + 'BPMNShape'):
        b = sh.find(DC + 'Bounds'); lb = sh.find(BD + 'BPMNLabel')
        lbb = lb.find(DC + 'Bounds') if lb is not None else None
        shapes[sh.get('bpmnElement')] = dict(x=float(b.get('x')), y=float(b.get('y')), w=float(b.get('width')), h=float(b.get('height')),
            label=(tuple(float(lbb.get(k)) for k in ('x', 'y', 'width', 'height')) if lbb is not None and float(lbb.get('width')) > 0 else None))
    for ed in root.iter(BD + 'BPMNEdge'):
        edges[ed.get('bpmnElement')] = [(float(w.get('x')), float(w.get('y'))) for w in ed.findall(DI_ + 'waypoint')]
    lanes = []
    for l in pr.iter(BM + 'lane'):
        g = shapes[l.get('id')]
        lanes.append(dict(id=l.get('id'), name=clean(l.get('name')), **g, refs=[r.text for r in l.findall(BM + 'flowNodeRef')]))
    KINDS = ('startEvent', 'endEvent', 'task', 'exclusiveGateway')
    nodes = []
    for e in pr:
        k = e.tag.replace(BM, '')
        if k not in KINDS: continue
        g = shapes[e.get('id')]
        doc = e.find(BM + 'documentation')
        n = dict(id=e.get('id'), name=clean(e.get('name')), kind=k, doc=clean(doc.text) if doc is not None and doc.text else '', gw=None, **g)
        cy = g['y'] + g['h'] / 2
        n['lane'] = next((l['id'] for l in lanes if n['id'] in l['refs']), None) or \
                    next((l['id'] for l in lanes if l['y'] <= cy <= l['y'] + l['h']), None)
        nodes.append(n)
    flows = [dict(id=f.get('id'), src=f.get('sourceRef'), dst=f.get('targetRef'), name=clean(f.get('name')), pts=edges[f.get('id')])
             for f in pr.findall(BM + 'sequenceFlow')]
    ids = {n['id'] for n in nodes}
    assert all(f['src'] in ids and f['dst'] in ids for f in flows), 'fluxo aponta para nó inexistente'
    notes = []
    for t in pr.findall(BM + 'textAnnotation'):
        tx = t.find(BM + 'text')
        notes.append(dict(id=t.get('id'), text=clean(tx.text if tx is not None else ''), **shapes[t.get('id')]))
    assocs = [dict(id=a.get('id'), src=a.get('sourceRef'), dst=a.get('targetRef'), pts=edges.get(a.get('id'), [])) for a in pr.findall(BM + 'association')]
    inc = {n['id']: [] for n in nodes}; out = {n['id']: [] for n in nodes}
    for f in flows: out[f['src']].append(f['id']); inc[f['dst']].append(f['id'])
    P = ['<?xml version="1.0" encoding="UTF-8"?>',
         '<bpmn:definitions xmlns:bpmn="http://www.omg.org/spec/BPMN/20100524/MODEL" xmlns:bpmndi="http://www.omg.org/spec/BPMN/20100524/DI" xmlns:dc="http://www.omg.org/spec/DD/20100524/DC" xmlns:di="http://www.omg.org/spec/DD/20100524/DI" id="Defs_%s" targetNamespace="http://skema.jr/rede-oftalmo">' % proc['id'],
         '<bpmn:collaboration id="Collab_1"><bpmn:participant id="Pool_1" name=%s processRef="Process_1" /></bpmn:collaboration>' % quoteattr(pool_name),
         '<bpmn:process id="Process_1" isExecutable="false"><bpmn:laneSet id="LaneSet_1">']
    for l in lanes:
        P.append('<bpmn:lane id="%s" name=%s>' % (l['id'], quoteattr(l['name'])))
        for n in nodes:
            if n['lane'] == l['id']: P.append('<bpmn:flowNodeRef>%s</bpmn:flowNodeRef>' % n['id'])
        P.append('</bpmn:lane>')
    P.append('</bpmn:laneSet>')
    for n in nodes:
        nm = (' name=%s' % quoteattr(n['name'])) if n['name'] else ''
        P.append('<bpmn:%s id="%s"%s>' % (n['kind'], n['id'], nm))
        if n['doc']: P.append('<bpmn:documentation>%s</bpmn:documentation>' % escape(n['doc']))
        for i in inc[n['id']]: P.append('<bpmn:incoming>%s</bpmn:incoming>' % i)
        for o in out[n['id']]: P.append('<bpmn:outgoing>%s</bpmn:outgoing>' % o)
        P.append('</bpmn:%s>' % n['kind'])
    for f in flows:
        nm = (' name=%s' % quoteattr(f['name'])) if f['name'] else ''
        P.append('<bpmn:sequenceFlow id="%s" sourceRef="%s" targetRef="%s"%s />' % (f['id'], f['src'], f['dst'], nm))
    for t in notes: P.append('<bpmn:textAnnotation id="%s"><bpmn:text>%s</bpmn:text></bpmn:textAnnotation>' % (t['id'], escape(t['text'])))
    for a in assocs: P.append('<bpmn:association id="%s" sourceRef="%s" targetRef="%s" />' % (a['id'], a['src'], a['dst']))
    P.append('</bpmn:process>')
    P.append('<bpmndi:BPMNDiagram id="Diagram_1"><bpmndi:BPMNPlane id="Plane_1" bpmnElement="Collab_1">')
    g = dict(shapes[pool_id])
    band = min(l['x'] for l in lanes) - g['x']          # faixa do título do pool; o export do Bizagi deixa só 10 px
    if band < 40: g['x'] -= 40 - band; g['w'] += 40 - band
    P.append('<bpmndi:BPMNShape id="Pool_1_di" bpmnElement="Pool_1" isHorizontal="true"><dc:Bounds x="%g" y="%g" width="%g" height="%g" /></bpmndi:BPMNShape>' % (g['x'], g['y'], g['w'], g['h']))
    for l in lanes:
        P.append('<bpmndi:BPMNShape id="%s_di" bpmnElement="%s" isHorizontal="true"><dc:Bounds x="%g" y="%g" width="%g" height="%g" /></bpmndi:BPMNShape>' % (l['id'], l['id'], l['x'], l['y'], l['w'], l['h']))
    for n in list(nodes) + notes:
        mk = ' isMarkerVisible="true"' if n.get('kind') == 'exclusiveGateway' else ''
        lab = ('<bpmndi:BPMNLabel><dc:Bounds x="%g" y="%g" width="%g" height="%g" /></bpmndi:BPMNLabel>' % n['label']) if n.get('label') and n.get('kind') != 'task' else ''
        P.append('<bpmndi:BPMNShape id="%s_di" bpmnElement="%s"%s><dc:Bounds x="%g" y="%g" width="%g" height="%g" />%s</bpmndi:BPMNShape>' % (n['id'], n['id'], mk, n['x'], n['y'], n['w'], n['h'], lab))
    for f in list(flows) + assocs:
        P.append('<bpmndi:BPMNEdge id="%s_di" bpmnElement="%s">' % (f['id'], f['id']))
        for (x, y) in f['pts']: P.append('<di:waypoint x="%g" y="%g" />' % (x, y))
        P.append('</bpmndi:BPMNEdge>')
    P.append('</bpmndi:BPMNPlane></bpmndi:BPMNDiagram></bpmn:definitions>')
    with open(os.path.join(OUT_B, proc['id'] + '.bpmn'), 'w', encoding='utf-8') as fh: fh.write('\n'.join(P))
    return dict(id=proc['id'], title=proc['title'], short=proc['short'], poolName=pool_name, sourceFile=proc['file'],
                bpmn='bpmn/%s.bpmn' % proc['id'],
                lanes=[dict(id=l['id'], name=l['name']) for l in lanes],
                nodes=[node_out(n, via_map(nodes, flows), seq_map(nodes, flows)) for n in nodes],
                flows=[dict(id=f['id'], src=f['src'], dst=f['dst'], name=f['name']) for f in flows],
                notes=[dict(id=t['id'], text=t['text']) for t in notes],
                counts=dict(tasks=sum(n['kind'] == 'task' for n in nodes), gateways=sum(n['kind'] == 'exclusiveGateway' for n in nodes),
                            events=sum(n['kind'].endswith('Event') for n in nodes), flows=len(flows), lanes=len(lanes)))

def convert_any(p):
    return convert_bpmn(p) if p['file'].endswith('.bpmn') else convert(p)

procs = [convert_any(p) for p in PROCS]

# origem de cada etapa do fluxo de córnea: desenho original da equipe × acrescentada a partir da entrevista
_ORIG = {'Node_entender', 'Node_fenda', 'Node_grau', 'Node_grins', 'Node_relatorio', 'Node_entrega'}
for _p in procs:
    for _n in _p['nodes']:
        _n['origin'] = ('diagrama' if _n['id'] in _ORIG else 'entrevista') if _p['id'] == 'exame-de-cornea' else 'diagrama'

# ---------------------------------------------------------------- planilha
# Medicoes_Rede_Oftalmo.xlsx: aba "1ºandar" (dados a partir da linha 3) e aba "2ºandar".
# Os tempos digitados como "2:17" ficam guardados como horário 02:17 — na planilha significam 2 min 17 s
# (hora = minutos, minuto = segundos). Início/Fim, ao contrário, são horários do relógio.
wb = openpyxl.load_workbook(os.path.join(SRC, 'Medicoes_Rede_Oftalmo.xlsx'), data_only=True)
W1, W2 = wb['1ºandar'], wb['2ºandar']
FIRST = 3
def mmss(v): return v.hour * 60 + v.minute if isinstance(v, datetime.time) else None
def tod(v): return v.hour * 3600 + v.minute * 60 + v.second if isinstance(v, datetime.time) else None
def hhmm(v): return v.strftime('%H:%M') if isinstance(v, datetime.time) else None
def cell(ws, col, r): return ws['%s%d' % (col, r)].value
def rows_of(ws): return range(FIRST, ws.max_row + 1)

steps = []
def add_step(id_, label, phase, samples, excluded=None, **kw):
    vals = [s['seconds'] for s in samples]
    steps.append(dict(id=id_, label=label, phase=phase, n=len(vals),
              mean=round(statistics.mean(vals), 2) if vals else None,
              median=statistics.median(vals) if vals else None,
              min=min(vals) if vals else None, max=max(vals) if vals else None,
              stdev=round(statistics.stdev(vals), 2) if len(vals) > 1 else None,
              samples=samples, excluded=excluded or [], **kw))

def direct(ws, id_, label, phase, ncol, tcol, ocol=None, kind='cronometrado', **kw):
    smp, exc = [], []
    for r in rows_of(ws):
        v = cell(ws, tcol, r)
        if v in (None, ''): continue
        sx = mmss(v)
        if sx is None: exc.append(dict(row=r, raw=str(v), reason='valor não é uma duração')); continue
        obs = cell(ws, ocol, r) if ocol else None
        smp.append(dict(row=r, obs=cell(ws, ncol, r) if ncol else None, seconds=sx, note=clean(obs) if obs else None))
    add_step(id_, label, phase, smp, exc, source='Colunas %s%s' % (ncol + '–' if ncol else '', tcol), kind=kind, **kw)

def start_end(ws, id_, label, phase, a, b, d):
    smp, exc = [], []
    for r in rows_of(ws):
        sa, sb = cell(ws, a, r), cell(ws, b, r)
        ta, tb = tod(sa), tod(sb)
        if ta is None and tb is None: continue
        if ta is None or tb is None: exc.append(dict(row=r, raw='%s – %s' % (sa, sb), reason='início ou fim ausente')); continue
        if tb - ta < 0: exc.append(dict(row=r, raw='%s – %s' % (hhmm(sa), hhmm(sb)), reason='horário de fim anterior ao de início')); continue
        smp.append(dict(row=r, obs=cell(ws, 'A', r), seconds=tb - ta, start=hhmm(sa), end=hhmm(sb), note=None))
    add_step(id_, label, phase, smp, exc, source='Colunas %s–%s' % (a, d), kind='início / fim')

# --- 1º andar: tempos cronometrados direto (mm:ss)
for id_, label, phase, cn, ct, co in [
  ('atendimento-glaucoma', 'Atendimento Glaucoma', 'Chegada', 'A', 'B', 'C'),
  ('microscopia', 'Microscopia', 'Exames', 'D', 'E', 'F'),
  ('topografia', 'Topografia', 'Exames', 'G', 'H', 'I'),
  ('tonometria-sopro', 'Tonometria de sopro', 'Exames', 'J', 'K', 'L'),
  ('auto-refratario', 'Auto refratário', 'Exames', 'M', 'N', 'O'),
  ('glaucoma-consulta', 'Glaucoma Consulta', 'Consulta', 'P', 'Q', 'R')]:
    direct(W1, id_, label, phase, cn, ct, co)
# --- Recepção (Portaria): aba 2ºandar, colunas A (nº) e B (tempo)
direct(W2, 'recepcao', 'Recepção (Portaria)', 'Chegada', 'A', 'B')
# a Recepção vem antes das demais etapas na lista
steps.insert(0, steps.pop())

# --- Campo visual: início/fim; nas linhas sem horário ("x") vale a duração cronometrada (coluna AA)
smp, exc = [], []
for r in rows_of(W1):
    a, b, d = cell(W1, 'Y', r), cell(W1, 'Z', r), cell(W1, 'AA', r)
    ta, tb = tod(a), tod(b)
    if ta is not None and tb is not None:
        calc = tb - ta
        if calc < 0: exc.append(dict(row=r, raw='%s – %s' % (hhmm(a), hhmm(b)), reason='horário de fim anterior ao de início')); continue
        item = dict(row=r, obs=cell(W1, 'A', r), seconds=calc, start=hhmm(a), end=hhmm(b), note=None)
        if isinstance(d, datetime.time) and mmss(d) != calc:
            item['flag'] = 'conferir'; item['note'] = 'A duração escrita na planilha difere do intervalo início–fim; usado o intervalo registrado.'
        smp.append(item)
    elif a == 'x' and isinstance(d, datetime.time):
        smp.append(dict(row=r, obs=cell(W1, 'A', r), seconds=mmss(d), note='Tempo cronometrado direto (sem horário de início/fim)'))
add_step('campo-visual', 'Campo visual', 'Exames', smp, exc, source='Colunas Y–AA', kind='cronometrado / início-fim')

start_end(W1, 'oct', 'OCT', 'Exames', 'S', 'T', 'U')
start_end(W1, 'paquimetria', 'Paquimetria corneana', 'Exames', 'V', 'W', 'X')
start_end(W1, 'retinografia', 'Retinografia', 'Exames', 'AB', 'AC', 'AD')

# --- Atendimento Córnea (antes "Pós-consulta"): tempo de atendimento (AH) + início/fim quando houver
smp, exc = [], []
for r in rows_of(W1):
    v = cell(W1, 'AH', r)
    if v in (None, ''): continue
    sx = mmss(v)
    if sx is None: exc.append(dict(row=r, raw=str(v), reason='valor não é uma duração')); continue
    a, b = cell(W1, 'AE', r), cell(W1, 'AF', r)
    item = dict(row=r, obs=cell(W1, 'A', r), seconds=sx, note=None)
    if tod(a) is not None and tod(b) is not None: item.update(start=hhmm(a), end=hhmm(b))
    smp.append(item)
add_step('atendimento-cornea', 'Atendimento Córnea', 'Córnea', smp, exc, source='Colunas AE–AF e AH', kind='cronometrado / início-fim')
# tempo sem paciente do posto da córnea (AG): só valores em mm:ss são válidos
smp, exc = [], []
for r in rows_of(W1):
    v = cell(W1, 'AG', r)
    if v in (None, ''): continue
    sx = mmss(v)
    if sx is None: exc.append(dict(row=r, raw=str(v) if v == 'x' else 'valor fora do formato', reason='sem valor ("x") ou fora do formato de tempo' )); continue
    smp.append(dict(row=r, obs=cell(W1, 'A', r), seconds=sx, note=None))
add_step('atendimento-cornea-ociosidade', 'Atendimento Córnea · tempo sem paciente', 'Córnea', smp, exc, source='Coluna AG', kind='intervalo do posto')

# --- 2º andar: Retina
direct(W2, 'retina-mario', 'Retina Dr. Mário Carlos', 'Retina', 'D', 'E', None)
smp = []
for col, nome in (('G', '1ª dilatação'), ('H', '2ª dilatação'), ('I', '3ª dilatação')):
    for r in rows_of(W2):
        v = cell(W2, col, r)
        if mmss(v) is not None: smp.append(dict(row=r, obs=nome, seconds=mmss(v), note=clean(cell(W2, 'J', r)) if cell(W2, 'J', r) else None))
add_step('dilatacao', 'Dilatação de colírio', 'Retina', smp, [], source='Colunas G–I', kind='cronometrado')
direct(W2, 'consultorio-retina', 'Consultório Retina', 'Retina', 'K', 'L', 'N')
direct(W2, 'ecobiometria', 'Ecobiometria / Ecografia', 'Retina', 'O', 'P', 'R')
direct(W2, 'atendimento-retina', 'Atendimento Retina', 'Retina', None, 'T', None)
RETINA = {'retina-mario', 'dilatacao', 'consultorio-retina', 'ecobiometria', 'atendimento-retina'}
empty_steps = []

# ---------------------------------------------------------------- associação etapa-do-fluxograma ↔ medição
g = next(p for p in procs if p['id'] == 'glaucoma')
cor = next(p for p in procs if p['id'] == 'exame-de-cornea')
def node(name_part):
    hits = [n for n in g['nodes'] if name_part.lower() in n['name'].lower() and n['kind'] == 'task']
    assert len(hits) == 1, (name_part, hits)
    return hits[0]['id']
lane_exames = next(l['id'] for l in g['lanes'] if l['name'] == 'Exames')
lane_medico_cor = next(l['id'] for l in cor['lanes'] if 'dico' in l['name'].lower())
LINK = {
  'recepcao': dict(process='glaucoma', nodes=[node('Portaria identifica'), node('Recepção abre a ficha')], confidence='provável',
       note='A planilha mede “Recepção (Portaria)” como uma etapa única; no fluxograma ela corresponde às duas tarefas da entrada (Portaria e Recepção). O mesmo tempo é exibido nas duas, sem soma.'),
  'auto-refratario': dict(process='glaucoma', nodes=[node('Refração')], confidence='provável', note='O auto refrator é o equipamento da etapa de refração.'),
  'tonometria-sopro': dict(process='glaucoma', nodes=[node('Mapeamento de')], confidence='provável',
       note='A tarefa do fluxograma une “mapeamento de retina e tonometria”; a planilha mede a tonometria de sopro separadamente.'),
  'retinografia': dict(process='glaucoma', nodes=[node('Mapeamento de')], confidence='provável',
       note='A tarefa do fluxograma une “mapeamento de retina e tonometria”; a planilha mede a retinografia separadamente.'),
}
LANE_STEPS = {lane_exames: ['microscopia', 'topografia', 'oct', 'paquimetria', 'campo-visual', 'glaucoma-consulta']}
for st in steps:
    sid = st['id']
    st['group'] = 'retina' if sid in RETINA else 'principal'
    if sid in RETINA:
        st['link'] = dict(process='retina', nodes=[], confidence='sem fluxograma', note='A Retina não tem fluxograma mapeado; as medições aparecem só nos resultados.')
    elif sid.startswith('atendimento-cornea'):
        st['link'] = dict(process='exame-de-cornea', nodes=[], lane=lane_medico_cor, confidence='raia',
                          note='A planilha mede o atendimento da córnea como um bloco, do início ao fim; no fluxograma ele corresponde à raia do médico (da conversa inicial à entrega da receita). Associação por raia.')
    elif sid in LINK: st['link'] = LINK[sid]
    elif sid in LANE_STEPS[lane_exames]:
        st['link'] = dict(process='glaucoma', nodes=[], lane=lane_exames, confidence='raia', note='O fluxograma não detalha este exame/etapa individualmente; a medição é associada à raia “Exames”.')
    else:
        st['link'] = dict(process='glaucoma', nodes=[], confidence='sem etapa equivalente', note='Não há tarefa com este nome no fluxograma; mantida apenas nos dados do processo.')
    st['process'] = st['link']['process']
    if st['kind'] == 'intervalo do posto':
        st['link'] = dict(process=st['link']['process'], nodes=[], confidence='sem etapa equivalente', note='Intervalo em que o posto ficou sem paciente; não é uma etapa da jornada.')

n_valid = sum(len(x['samples']) for x in steps if x['kind'] != 'intervalo do posto')
quality = [
 dict(id='fonte', level='info', title='Fonte dos dados', text='Medicoes_Rede_Oftalmo.xlsx (abas 1ºandar e 2ºandar). Tempos digitados como “2:17” foram lidos como 2 min 17 s.'),
 dict(id='independentes', level='info', title='Medições independentes por etapa', text='Cada etapa tem a sua própria contagem de pacientes; por isso o site não soma etapas nem calcula tempo total de permanência.'),
 dict(id='excluidos', level='corrigido', title='Valores desconsiderados', text='Valores que não são durações válidas (“x”, horários fora do formato) foram desconsiderados e estão listados em cada etapa. Nenhum valor válido foi alterado.'),
]

# ---------------------------------------------------------------- agenda de horários × cronoanálise (sem contar 2 vezes)
import math
ag_ws = openpyxl.load_workbook(os.path.join(SRC, 'Agenda_horarios.xlsx'), data_only=True)['Horários']
ag_rows = [r for r in ag_ws.iter_rows(min_row=2, values_only=True) if r[0]]
hm = lambda t: t.strftime('%H:%M') if hasattr(t, 'strftime') else None
AG_COLS = {'campo-visual': (4, 5), 'retinografia': (6, 7), 'paquimetria': (8, 9), 'oct': (10, 11)}
have = {(st['id'], x.get('start'), x.get('end')) for st in steps for x in st['samples'] if x.get('start')}
ag_total = ag_match = 0
for r in ag_rows:
    for sid, (a, b) in AG_COLS.items():
        if r[a] and r[b]:
            ag_total += 1; ag_match += (sid, hm(r[a]), hm(r[b])) in have
ag_dates = sorted({datetime.datetime.strptime(r[0], '%d/%m/%Y').date() for r in ag_rows})
print('agenda: %d pacientes, %d tempos de exame, %d já na cronoanálise' % (len(ag_rows), ag_total, ag_match))
quality.append(dict(id='agenda', level='info', title='Agenda de horários já está na cronoanálise',
    text='A planilha “agenda de horários” (%d pacientes em %s) traz Campo visual, Retinografia, Paquimetria e OCT. Dos %d tempos de exame dela, %d já constam na cronoanálise (mesmo início e fim) e não foram somados de novo: o site usa só a cronoanálise.' % (
        len(ag_rows), ' e '.join(d.strftime('%d/%m') for d in ag_dates), ag_total, ag_match)))

# ---------------------------------------------------------------- cronograma de coleta
PLAN_TARGET, PLAN_PER_VISIT, PLAN_PER_WEEK = 100, 35, 2
PLAN_EXCLUDED = ['oct', 'campo-visual', 'retinografia', 'paquimetria']
plan_steps = [st for st in steps if st['id'] not in PLAN_EXCLUDED and st['group'] == 'principal' and st['kind'] != 'intervalo do posto']
rows_ = [dict(id=st['id'], label=st['label'], n=st['n'], missing=max(0, PLAN_TARGET - st['n']),
              visits=math.ceil(max(0, PLAN_TARGET - st['n']) / PLAN_PER_VISIT)) for st in plan_steps]
n_visits = max(r['visits'] for r in rows_)
weekdays = sorted({d.weekday() for d in ag_dates})                  # dias da semana das visitas já feitas (qui/sex)
d = ag_dates[-1] + datetime.timedelta(days=1); visits_dates = []
while len(visits_dates) < n_visits:
    if d.weekday() in weekdays: visits_dates.append(d)
    d += datetime.timedelta(days=1)
def after(i, r): return min(PLAN_TARGET, r['n'] + PLAN_PER_VISIT * i)
visits_ = []
for i, dt in enumerate(visits_dates, 1):
    visits_.append(dict(n=i, date=dt.isoformat(), closes=[r['label'] for r in rows_ if r['visits'] == i],
                        lowest=min(after(i, r) for r in rows_)))
nxt = d
while nxt.weekday() not in weekdays: nxt += datetime.timedelta(days=1)
# início do calendário: próximo dia de coleta (qui/sex) depois da data em que o site é gerado
start = datetime.date.today() + datetime.timedelta(days=1)
while start.weekday() not in weekdays: start += datetime.timedelta(days=1)
plan = dict(target=PLAN_TARGET, perVisit=PLAN_PER_VISIT, perWeek=PLAN_PER_WEEK, excluded=[st['label'] for st in steps if st['id'] in PLAN_EXCLUDED],
            steps=rows_, visits=n_visits, schedule=visits_, lastDone=ag_dates[-1].isoformat(), end=visits_dates[-1].isoformat(), next=nxt.isoformat(),
            bottleneck=[r['label'] for r in rows_ if r['visits'] == n_visits],
            weekdays=weekdays, start=start.isoformat())
print('plano:', n_visits, 'visitas →', plan['end'], [r['label'] + ' ' + str(r['visits']) for r in rows_])

# ---------------------------------------------------------------- rótulos de tempo total dentro dos fluxogramas
for p in procs:
    p['totalLabel'] = None
    for n in p['nodes'] + [dict(name=t['text']) for t in p.get('notes', [])]:
        m = re.search(r'tempo total aprox\.?\s*([0-9]+:[0-9]{2})(.*)$', n['name'])
        if m: p['totalLabel'] = dict(value=m.group(1), note=clean(m.group(2)).strip(' ()'), text=n['name'])
        m2 = re.search(r'só da etapa de consulta.*aprox\.\s*([0-9]+:[0-9]{2})', n['name'])
        if m2: p['totalLabel'] = dict(value=m2.group(1), note='só a etapa de consulta', text=n['name'])

# ---------------------------------------------------------------- saída
json.dump(dict(processes=procs), open(os.path.join(OUT_D, 'processes.json'), 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
json.dump(dict(source='Medicoes_Rede_Oftalmo.xlsx · abas “1ºandar” e “2ºandar”', unit='segundos', steps=steps,
               plan=plan, emptySteps=empty_steps, laneSteps={k: v for k, v in LANE_STEPS.items()}, quality=quality),
          open(os.path.join(OUT_D, 'measurements.json'), 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
for s in steps: print('%-34s n=%-3d mean=%-8s med=%-6s min=%-5s max=%-6s exc=%d' % (s['label'], s['n'], s['mean'], s['median'], s['min'], s['max'], len(s['excluded'])))
for p in procs: print(p['id'], p['counts'], p['totalLabel'])
